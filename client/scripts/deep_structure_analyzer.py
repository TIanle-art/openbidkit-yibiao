import docx
import pymupdf
import sys
from docx.oxml import parse_xml

sys.stdout.reconfigure(encoding='utf-8')

def audit_document_full(pdf_path, docx_path, label):
    print(f"\n=======================================================")
    print(f"       全维度深度审计分析: {label}")
    print(f"=======================================================")
    
    pdf = pymupdf.open(pdf_path)
    doc = docx.Document(docx_path)
    
    print(f"1. 页面规模与分节:")
    print(f"   - PDF 总页数: {len(pdf)}")
    print(f"   - DOCX 分节数 (Sections): {len(doc.sections)}")
    print(f"   - DOCX 段落总数: {len(doc.paragraphs)}")
    print(f"   - DOCX 表格总数: {len(doc.tables)}")
    print(f"   - DOCX 内嵌图形/图片数: {len(doc.inline_shapes)}")
    
    # Check sections margins
    margins_set = set()
    for i, sec in enumerate(doc.sections):
        top = round(sec.top_margin.pt, 1)
        bot = round(sec.bottom_margin.pt, 1)
        left = round(sec.left_margin.pt, 1)
        right = round(sec.right_margin.pt, 1)
        margins_set.add((top, bot, left, right))
    print(f"   - 页边距组合数: {len(margins_set)} (抽样: {list(margins_set)[:3]})")
    
    # Check tables layout
    autofit_count = 0
    fixed_count = 0
    for t in doc.tables:
        tblPr = t._tbl.tblPr
        tblLayout = tblPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblLayout')
        if tblLayout is not None:
            val = tblLayout.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}type')
            if val == 'fixed':
                fixed_count += 1
            else:
                autofit_count += 1
        else:
            autofit_count += 1
    print(f"2. 表格布局锁定状态:")
    print(f"   - 固定网格锁定 (fixed): {fixed_count} 个")
    print(f"   - 浮动自适应 (autofit，易跑版): {autofit_count} 个")
    
    # Check images aspect ratio & size
    print(f"3. 图片几何与画质审计 (前 10 张图):")
    for i, shape in enumerate(list(doc.inline_shapes)[:10]):
        w_in = round(shape.width.inches, 2)
        h_in = round(shape.height.inches, 2)
        print(f"   - 图片 [{i+1}]: 宽={w_in} in, 高={h_in} in, 比例={round(w_in/h_in, 2) if h_in else 'N/A'}")
        
    # Check Fonts in Paragraphs
    font_counter = {}
    snap_to_grid_count = 0
    exact_line_count = 0
    for p in doc.paragraphs:
        pPr = p._p.get_or_add_pPr()
        snap = pPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}snapToGrid')
        if snap is not None and snap.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val') == '0':
            snap_to_grid_count += 1
            
        spacing = pPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}spacing')
        if spacing is not None and spacing.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}lineRule') == 'exact':
            exact_line_count += 1
            
        for r in p.runs:
            fn = r.font.name or "未定义"
            font_counter[fn] = font_counter.get(fn, 0) + 1
            
    print(f"4. 格式稳定性指标:")
    print(f"   - 关闭对齐网格 (snapToGrid=0): {snap_to_grid_count}/{len(doc.paragraphs)} ({round(snap_to_grid_count/max(1,len(doc.paragraphs))*100, 1)}%)")
    print(f"   - 锁定固定行距 (lineRule=exact): {exact_line_count}/{len(doc.paragraphs)} ({round(exact_line_count/max(1,len(doc.paragraphs))*100, 1)}%)")
    print(f"   - 字体分布: {font_counter}")

if __name__ == "__main__":
    audit_document_full(
        r"d:/APPs/Downloads/文件/赛鼎 技术文件0818(1).pdf",
        r"d:/APPs/Downloads/文件/已生成标书成果/01-赛鼎起重机_原版复刻_优化版.docx",
        "赛鼎项目 (Project 1)"
    )
    audit_document_full(
        r"d:/APPs/Downloads/文件/圣起过轨吊投标文件.pdf",
        r"d:/APPs/Downloads/文件/已生成标书成果/02-磨煤机过轨吊_原版复刻_优化版.docx",
        "磨煤机项目 (Project 2)"
    )
