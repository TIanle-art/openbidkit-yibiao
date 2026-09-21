import docx
import sys
import zipfile

sys.stdout.reconfigure(encoding='utf-8')

def audit_file(docx_path, label):
    print(f"\n=======================================================")
    print(f"  Flagship Verification Audit: {label}")
    print(f"  File: {docx_path}")
    print(f"=======================================================")
    doc = docx.Document(docx_path)
    
    # 1. Rogue check
    rogue_paras = []
    for i, p in enumerate(doc.paragraphs[10:]):
        t = p.text.strip()
        if any(k in t for k in ["河南省中原奥起实业有限公司投标文件", 
                                "内蒙古长城发电有限公司上海庙煤电项目中速磨煤机用过轨起重机设备采购", 
                                "※投标单位：河南圣起机械集团有限公司"]):
            rogue_paras.append((i+10, t[:40]))
    print(f"1. 残留正文页眉/页脚段落: {len(rogue_paras)} 处 (目标: 0)")
    
    # 2. Snap to grid & Line rule
    snap_count = 0
    exact_line_count = 0
    east_asia_runs = 0
    for p in doc.paragraphs:
        pPr = p._p.get_or_add_pPr()
        snap = pPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}snapToGrid')
        if snap is not None and snap.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val') == '0':
            snap_count += 1
            
        spacing = pPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}spacing')
        if spacing is not None and spacing.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}lineRule') == 'exact':
            exact_line_count += 1
            
        for r in p.runs:
            rPr = r._r.get_or_add_rPr()
            rFonts = rPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rFonts')
            if rFonts is not None and rFonts.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}hint') == 'eastAsia':
                east_asia_runs += 1
                
    total_p = len(doc.paragraphs)
    print(f"2. 格式锁定指标:")
    print(f"   - 关闭对齐网格 (snapToGrid=0): {snap_count}/{total_p} ({round(snap_count/max(1,total_p)*100, 1)}%)")
    print(f"   - 锁定固定行距 (lineRule=exact): {exact_line_count}/{total_p} ({round(exact_line_count/max(1,total_p)*100, 1)}%)")
    print(f"   - 注入东亚中文字体标识 (hint=eastAsia): {east_asia_runs} 处")
    
    # 3. Table fixed layout
    fixed_tbls = 0
    for t in doc.tables:
        tblPr = t._tbl.tblPr
        tblLayout = tblPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblLayout')
        if tblLayout is not None and tblLayout.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}type') == 'fixed':
            fixed_tbls += 1
    print(f"3. 表格定宽锁定: {fixed_tbls}/{len(doc.tables)} ({round(fixed_tbls/max(1,len(doc.tables))*100, 1)}%)")
    
    # 4. doNotCompressPictures check
    with zipfile.ZipFile(docx_path, 'r') as z:
        settings_xml = z.read("word/settings.xml").decode('utf-8')
        has_nocompress = "<w:doNotCompressPictures" in settings_xml
    print(f"4. 画质无损锁 (doNotCompressPictures): {'[已激活]' if has_nocompress else '[未激活]'}")

if __name__ == "__main__":
    audit_file(r"d:\APPs\Downloads\文件\已生成标书成果\01-赛鼎起重机_原版复刻_完美修复版.docx", "赛鼎完美修复版")
    audit_file(r"d:\APPs\Downloads\文件\已生成标书成果\02-磨煤机过轨吊_原版复刻_完美修复版.docx", "磨煤机完美修复版")
