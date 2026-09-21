import os
import sys
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

def apply_master_visual_polish(docx_path, project_type):
    print(f"\n[MASTER-POLISH] Starting visual & font polish for {project_type}: {docx_path}")
    doc = docx.Document(docx_path)
    
    # 1. Determine Header & Logo Config
    logo_path = ""
    header_text = ""
    footer_text = ""
    logo_w_in = 0.45
    logo_h_in = 0.45

    if project_type == "sading":
        logo_path = r"d:\APPs\Downloads\文件\sading_logo.png"
        header_text = "河南省中原奥起实业有限公司投标文件"
        logo_w_in = 0.45
        logo_h_in = 0.45
    elif project_type == "guagui":
        logo_path = r"d:\APPs\Downloads\文件\shengqi_logo.png"
        header_text = "内蒙古长城发电有限公司上海庙煤电项目中速磨煤机用过轨起重机设备采购 投标文件"
        footer_text = "※投标单位：河南圣起机械集团有限公司 创新无限，信誉永恒！"
        logo_w_in = 0.55
        logo_h_in = 0.45

    # 2. Inject Authentic Header Table (Text on Left, Logo on Right, Divider on Bottom)
    for s_idx, sec in enumerate(doc.sections):
        # Clear existing header paragraphs
        hdr = sec.header
        for hp in hdr.paragraphs:
            hp.text = ""
        # Remove any existing header tables
        for ht in hdr.tables:
            ht._tbl.getparent().remove(ht._tbl)

        # Create 2-column header table
        # Table width ~ 6.5 inches (standard printable width)
        htbl = hdr.add_table(rows=1, cols=2, width=Inches(6.5))
        htbl.alignment = docx.enum.table.WD_TABLE_ALIGNMENT.CENTER
        
        # Bottom border only
        htblBorders = parse_xml(r'''
            <w:tblBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
                <w:top w:val="none"/>
                <w:left w:val="none"/>
                <w:bottom w:val="single" w:sz="8" w:space="0" w:color="999999"/>
                <w:right w:val="none"/>
                <w:insideH w:val="none"/>
                <w:insideV w:val="none"/>
            </w:tblBorders>
        ''')
        htbl._tbl.tblPr.append(htblBorders)

        # Left Cell: Header Text
        c_left = htbl.cell(0, 0)
        c_left.width = Inches(5.5)
        pl = c_left.paragraphs[0]
        pl.alignment = WD_ALIGN_PARAGRAPH.LEFT
        rl = pl.add_run(header_text)
        rl_rPr = rl._r.get_or_add_rPr()
        rl_fonts = parse_xml(r'<w:rFonts xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:eastAsia="宋体" w:cs="Times New Roman" w:hint="eastAsia"/>')
        rl_rPr.append(rl_fonts)
        rl_sz = parse_xml(r'<w:sz xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:val="18"/>')
        rl_rPr.append(rl_sz)
        rl_col = parse_xml(r'<w:color xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:val="333333"/>')
        rl_rPr.append(rl_col)

        # Right Cell: Logo Image
        c_right = htbl.cell(0, 1)
        c_right.width = Inches(1.0)
        pr = c_right.paragraphs[0]
        pr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if os.path.exists(logo_path):
            rr = pr.add_run()
            rr.add_picture(logo_path, width=Inches(logo_w_in), height=Inches(logo_h_in))

        # Footer configuration
        if footer_text:
            ftr = sec.footer
            if len(ftr.paragraphs) > 0:
                fp = ftr.paragraphs[0]
                fp.text = ""
                fr = fp.add_run(footer_text)
                fr_rPr = fr._r.get_or_add_rPr()
                fr_fonts = parse_xml(r'<w:rFonts xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:eastAsia="宋体" w:cs="Times New Roman" w:hint="eastAsia"/>')
                fr_rPr.append(fr_fonts)
                fr_sz = parse_xml(r'<w:sz xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:val="18"/>')
                fr_rPr.append(fr_sz)
                fr_col = parse_xml(r'<w:color xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:val="666666"/>')
                fr_rPr.append(fr_col)

    # 3. Clean Rogue Logos and Polish Paragraph Fonts in Body
    removed_rogue_images = 0
    styled_headings = 0
    styled_body_paras = 0

    for p_idx, p in enumerate(doc.paragraphs):
        t = p.text.strip()

        # Check and remove rogue logo images accidentally left in body
        for r in p.runs:
            drawings = r._r.findall(qn('w:drawing'))
            for d in drawings:
                extent = d.find(qn('wp:extent'))
                if extent is not None:
                    cx = int(extent.get('cx', 0))
                    cy = int(extent.get('cy', 0))
                    w_pt = cx / 12700
                    h_pt = cy / 12700
                    # Rogue header logos are small square icons (around 25-50pt)
                    if 20 < w_pt < 60 and 20 < h_pt < 60 and (t == "" or header_text in t or "投标文件" in t):
                        r._r.remove(d)
                        removed_rogue_images += 1

        # Classify and Style Headings vs Body
        is_l1_heading = any(t.startswith(k) for k in [
            "1．", "2．", "3．", "4.", "5.", "6.", "7.", "8.", "9.", "10.", "11.", "12.", "13.",
            "一、", "二、", "三、", "四、", "五、", "六、", "七、", "八、", "九、", "十、", "十二、",
            "技术响应文件", "投 标 文 件"
        ]) and len(t) < 60 and not ("。" in t or "；" in t)

        is_l2_heading = any(t.startswith(k) for k in [
            "1.1", "1.2", "2.1", "2.2", "3.1", "3.2", "3.3", "3.4", "3.5", "3.6", "3.7",
            "4.1", "4.2", "5.1", "5.2", "（一）", "（二）", "（三）", "（四）", "（五）", "附件"
        ]) and len(t) < 60 and not ("。" in t or "；" in t)

        # Style runs inside paragraph
        for r in p.runs:
            rPr = r._r.get_or_add_rPr()
            
            # Determine font family
            target_eastasia = "宋体"
            target_sz = 12.0 # 小四
            is_bold = r.bold

            if is_l1_heading:
                target_eastasia = "黑体"
                target_sz = 16.0 # 三号
                is_bold = True
                styled_headings += 1
            elif is_l2_heading:
                target_eastasia = "黑体"
                target_sz = 14.0 # 四号
                is_bold = True
                styled_headings += 1
            elif "\t" in p.text and ("响应函" in p.text or "投标函" in p.text or "资格审查" in p.text):
                # TOC line
                target_eastasia = "黑体" if is_bold else "宋体"
                target_sz = 10.5 # 五号
            else:
                styled_body_paras += 1
                # Standard body font: 12pt SimSun
                target_eastasia = "宋体"
                # Preserve existing explicit heading size if already larger
                sz_elem = rPr.find(qn('w:sz'))
                if sz_elem is not None and sz_elem.get(qn('w:val')):
                    val = float(sz_elem.get(qn('w:val'))) / 2.0
                    if val >= 20:
                        target_sz = val

            # Set w:rFonts with w:hint="eastAsia" so Word displays proper font box
            rFonts = rPr.find(qn('w:rFonts'))
            if rFonts is None:
                rFonts = parse_xml(f'<w:rFonts xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:eastAsia="{target_eastasia}" w:cs="Times New Roman" w:hint="eastAsia"/>')
                rPr.append(rFonts)
            else:
                rFonts.set(qn('w:ascii'), "Times New Roman")
                rFonts.set(qn('w:hAnsi'), "Times New Roman")
                rFonts.set(qn('w:eastAsia'), target_eastasia)
                rFonts.set(qn('w:cs'), "Times New Roman")
                rFonts.set(qn('w:hint'), "eastAsia")

            # Set size
            val_half_pt = str(int(target_sz * 2))
            sz = rPr.find(qn('w:sz'))
            if sz is not None:
                sz.set(qn('w:val'), val_half_pt)
            else:
                rPr.append(parse_xml(f'<w:sz xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:val="{val_half_pt}"/>'))
            szCs = rPr.find(qn('w:szCs'))
            if szCs is not None:
                szCs.set(qn('w:val'), val_half_pt)
            else:
                rPr.append(parse_xml(f'<w:szCs xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:val="{val_half_pt}"/>'))

            # Set bold if heading
            if is_bold:
                r.bold = True

            # Ensure pure black color
            col = rPr.find(qn('w:color'))
            if col is not None:
                col.set(qn('w:val'), "000000")
            else:
                rPr.append(parse_xml(r'<w:color xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:val="000000"/>'))

    # 4. Standardize Tables
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for cp in cell.paragraphs:
                    for cr in cp.runs:
                        crPr = cr._r.get_or_add_rPr()
                        crFonts = crPr.find(qn('w:rFonts'))
                        if crFonts is not None:
                            crFonts.set(qn('w:eastAsia'), "宋体")
                            crFonts.set(qn('w:ascii'), "Times New Roman")
                            crFonts.set(qn('w:hint'), "eastAsia")
                        csz = crPr.find(qn('w:sz'))
                        if csz is not None:
                            csz.set(qn('w:val'), "21") # 10.5 pt
                        cszCs = crPr.find(qn('w:szCs'))
                        if cszCs is not None:
                            cszCs.set(qn('w:val'), "21")

    # Safe Save with Lock Handling
    save_path = docx_path
    try:
        doc.save(docx_path)
    except PermissionError:
        base_dir = os.path.dirname(docx_path)
        base_name = os.path.splitext(os.path.basename(docx_path))[0]
        save_path = os.path.join(base_dir, f"{base_name}_优化版.docx")
        print(f"  [锁定告警] 目标文件被 Word 占用，自动安全导出为: {save_path}")
        doc.save(save_path)

    print(f"  [完成] 页眉双栏表格与高清 Logo 装配成功！")
    print(f"  [完成] 移除了正文中的 {removed_rogue_images} 个散落游离 Logo")
    print(f"  [完成] 统一了全篇东亚字体（宋体/黑体带 hint=eastAsia，英数 Times New Roman）")
    print(f"  [完成] 成功保存: {save_path}")
    return save_path

if __name__ == "__main__":
    targets = [
        (r"d:\APPs\Downloads\文件\已生成标书成果\01-赛鼎起重机_AI生成.docx", "sading"),
        (r"d:\APPs\Downloads\文件\已生成标书成果\01-赛鼎起重机_原版复刻.docx", "sading"),
        (r"d:\APPs\Downloads\文件\已生成标书成果\02-磨煤机过轨吊_AI生成.docx", "guagui"),
        (r"d:\APPs\Downloads\文件\已生成标书成果\02-磨煤机过轨吊_原版复刻.docx", "guagui")
    ]
    for path, p_type in targets:
        apply_master_visual_polish(path, p_type)
