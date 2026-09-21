import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

docx_path = r"d:\APPs\Downloads\文件\已生成标书成果\02-磨煤机过轨吊_原版复刻_完美修复版.docx"
d = docx.Document(docx_path)

print("Searching for paragraphs near （二）近年财务状况表:")
for i, p in enumerate(d.paragraphs):
    if "近年财务状况表" in p.text:
        print(f"Found heading at P[{i}]: \"{p.text}\"")
        for j in range(max(0, i-2), min(len(d.paragraphs), i+15)):
            pj = d.paragraphs[j]
            has_drawing = len(pj._p.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing'))
            print(f"  P[{j}]: text=\"{pj.text.strip()[:40]}\", drawings={has_drawing}, pPr_xml={pj._p.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pPr') is not None}")
            if has_drawing:
                print(f"    P[{j}] XML:\n{pj._p.xml[:500]}")
        break
