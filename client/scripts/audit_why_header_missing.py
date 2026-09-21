import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

def audit_doc(docx_path, name):
    print(f"\n=======================================================")
    print(f"  Audit: {name}")
    print(f"  File: {docx_path}")
    print(f"=======================================================")
    doc = docx.Document(docx_path)
    
    # 1. Header audit across sections
    total_secs = len(doc.sections)
    header_counts = 0
    empty_header_secs = []
    sample_headers = []
    
    for s_idx, sec in enumerate(doc.sections):
        hdr = sec.header
        has_content = False
        tbl_count = len(hdr.tables)
        para_texts = [p.text.strip() for p in hdr.paragraphs if p.text.strip()]
        
        if tbl_count > 0:
            has_content = True
            tbl_text = [cell.text.strip() for row in hdr.tables[0].rows for cell in row.cells]
        else:
            tbl_text = []
            
        if para_texts or tbl_text:
            has_content = True
            header_counts += 1
            if len(sample_headers) < 5:
                sample_headers.append((s_idx, para_texts, tbl_text))
        else:
            if not sec.header.is_linked_to_previous:
                empty_header_secs.append(s_idx)
                
    print(f"1. 页眉覆盖审计:")
    print(f"   - 总分节数 (Sections): {total_secs}")
    print(f"   - 显式含页眉分节数: {header_counts}")
    print(f"   - 未链接且空页眉分节数: {len(empty_header_secs)} 处 (如分节: {empty_header_secs[:10]})")
    print(f"   - 页眉抽样: {sample_headers}")
    
    # 2. Images count audit
    total_inline = len(list(doc.inline_shapes))
    drawing_count = 0
    for p in doc.paragraphs:
        for r in p.runs:
            d = r._r.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing')
            drawing_count += len(d)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        d = r._r.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing')
                        drawing_count += len(d)
    print(f"\n2. 图片总数审计:")
    print(f"   - doc.inline_shapes 数量: {total_inline}")
    print(f"   - XML drawing 总数: {drawing_count}")

audit_doc(r"d:\APPs\Downloads\文件\已生成标书成果\01-赛鼎起重机_原版复刻_终极保真版.docx", "赛鼎终极保真版")
audit_doc(r"d:\APPs\Downloads\文件\已生成标书成果\02-磨煤机过轨吊_原版复刻_终极保真版.docx", "磨煤机终极保真版")
