import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

docx_path = r"d:\APPs\Downloads\文件\已生成标书成果\02-磨煤机过轨吊_原版复刻_完美修复版.docx"
d = docx.Document(docx_path)

for i, p in enumerate(d.paragraphs[73:], start=73):
    if "财务状况表" in p.text:
        print(f"Found body heading at P[{i}]: \"{p.text}\"")
        for j in range(max(73, i-1), min(len(d.paragraphs), i+6)):
            pj = d.paragraphs[j]
            has_drawing = len(pj._p.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing'))
            print(f"  P[{j}]: text=\"{pj.text.strip()[:40]}\", drawings={has_drawing}")
            if has_drawing:
                print(f"    P[{j}] XML:\n{pj._p.xml[:700]}")
        break
