import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

def audit_text_errors(docx_path, name):
    print(f"\n=======================================================")
    print(f"  Auditing Text Formatting Errors: {name}")
    print(f"  File: {docx_path}")
    print(f"=======================================================")
    d = docx.Document(docx_path)
    
    issues = []
    
    # 1. Check font sizes in body
    font_sizes = {}
    for i, p in enumerate(d.paragraphs):
        t = p.text.strip()
        if not t:
            continue
        pPr = p._p.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pPr')
        # Check alignment
        align = p.alignment
        # Check runs
        for r in p.runs:
            sz = r.font.size.pt if r.font.size else "default"
            font_sizes[sz] = font_sizes.get(sz, 0) + 1
            
        # Check if line contains weird characters or strange tabs
        if "\t\t\t" in p.text:
            issues.append((i, "Excessive tabs", t[:40]))
        if any(bad in t for bad in ["undefined", "NaN", "null", "TODO", "???"]):
            issues.append((i, "Suspicious placeholder text", t[:40]))
            
    print("Font size distribution in document runs:")
    for sz, count in sorted(font_sizes.items(), key=lambda x: str(x[0])):
        print(f"  {sz} pt: {count} runs")
        
    print(f"\nPotential Text Issues detected: {len(issues)}")
    for idx, kind, sample in issues[:15]:
        print(f"  P[{idx}]: [{kind}] \"{sample}\"")
        
    # Check sample paragraphs from various sections
    print("\nSample Content Paragraphs:")
    for i in [15, 20, 30, 50, 70, 90, 120, 150, 200]:
        if i < len(d.paragraphs):
            p = d.paragraphs[i]
            if p.text.strip():
                print(f"  P[{i}]: align={p.alignment}, text=\"{p.text.strip()[:60]}\"")

audit_text_errors(r"d:\APPs\Downloads\文件\已生成标书成果\test_sading_purged.docx", "赛鼎")
audit_text_errors(r"d:\APPs\Downloads\文件\已生成标书成果\test_guagui_purged.docx", "磨煤机")

