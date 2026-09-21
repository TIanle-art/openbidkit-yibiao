import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

docx_path = r"d:\APPs\Downloads\文件\已生成标书成果\02-磨煤机过轨吊_原版复刻_完美修复版.docx"
d = docx.Document(docx_path)

print("Searching for floating images with <wp:anchor>:")
anchors_found = 0
for i, p in enumerate(d.paragraphs):
    anchors = p._p.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}anchor')
    if anchors:
        anchors_found += len(anchors)
        if anchors_found <= 10:
            for a in anchors:
                posV = a.find('{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}positionV')
                pos_str = ""
                if posV is not None:
                    offset = posV.find('{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}posOffset')
                    if offset is not None:
                        pos_str = f"posOffset={offset.text}"
                extent = a.find('{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent')
                ext_str = f"cx={extent.get('cx')}, cy={extent.get('cy')}" if extent is not None else ""
                print(f"P[{i}]: anchor ({pos_str}, {ext_str})")

print(f"Total <wp:anchor> floating shapes found: {anchors_found}")

inline_found = 0
for i, p in enumerate(d.paragraphs):
    inlines = p._p.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}inline')
    if inlines:
        inline_found += len(inlines)

print(f"Total <wp:inline> inline shapes found: {inline_found}")
