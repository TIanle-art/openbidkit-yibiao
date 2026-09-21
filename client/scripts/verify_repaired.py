import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

def audit_repaired(docx_path, name):
    print(f"\n=======================================================")
    print(f"  Verification: {name}")
    print(f"  File: {docx_path}")
    print(f"=======================================================")
    doc = docx.Document(docx_path)
    
    # Header check
    total_secs = len(doc.sections)
    s0_tbls = len(doc.sections[0].header.tables)
    s0_text = [c.text.strip() for r in doc.sections[0].header.tables[0].rows for c in r.cells] if s0_tbls > 0 else []
    linked_secs = sum(1 for s in doc.sections[1:] if s.header.is_linked_to_previous)
    
    print(f"1. 页眉覆盖验证:")
    print(f"   - 总分节数: {total_secs}")
    print(f"   - Section 0 页眉表格: {s0_tbls} 个 (内容: {s0_text})")
    print(f"   - 继承 Section 0 页眉的分节数: {linked_secs}/{total_secs-1} ({round(linked_secs/(total_secs-1)*100, 1)}%)")
    
    # Image clipping check
    clipped_images = 0
    for p in doc.paragraphs:
        if p._p.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing'):
            pPr = p._p.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pPr')
            if pPr is not None:
                sp = pPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}spacing')
                if sp is not None and sp.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}lineRule') == 'exact':
                    clipped_images += 1
    print(f"\n2. 图片被固定行距截断裁切数量: {clipped_images} 处 (目标: 0)")
    
    # Oversized image check
    oversized = 0
    for s in doc.inline_shapes:
        if s.width.inches > 6.3 or s.height.inches > 8.7:
            oversized += 1
    print(f"3. 超出页面边界溢出图片数量: {oversized} 处 (目标: 0)")

audit_repaired(r"d:\APPs\Downloads\文件\已生成标书成果\01-赛鼎起重机_原版复刻_完美修复版.docx", "赛鼎完美修复版")
audit_repaired(r"d:\APPs\Downloads\文件\已生成标书成果\02-磨煤机过轨吊_原版复刻_完美修复版.docx", "磨煤机完美修复版")
