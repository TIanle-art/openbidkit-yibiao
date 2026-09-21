import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

def check_first_sections(docx_path, name):
    d = docx.Document(docx_path)
    print(f"\n=== {name} First 5 Sections ===")
    for i in range(min(5, len(d.sections))):
        s = d.sections[i]
        print(f"Section [{i}]:")
        print(f"  start_type: {s.start_type}")
        print(f"  page_w: {round(s.page_width.inches, 2)} in, page_h: {round(s.page_height.inches, 2)} in")
        print(f"  margins: top={round(s.top_margin.inches, 2)}, bot={round(s.bottom_margin.inches, 2)}, left={round(s.left_margin.inches, 2)}, right={round(s.right_margin.inches, 2)}")
        print(f"  diff_first_page: {s.different_first_page_header_footer}")
        print(f"  header paras: {[p.text for p in s.header.paragraphs if p.text.strip()]}")
        print(f"  header tables: {len(s.header.tables)}")
        print(f"  header is_linked: {s.header.is_linked_to_previous}")

check_first_sections(r"d:\APPs\Downloads\文件\已生成标书成果\01-赛鼎起重机_原版复刻_终极保真版.docx", "赛鼎终极保真版")
check_first_sections(r"d:\APPs\Downloads\文件\已生成标书成果\02-磨煤机过轨吊_原版复刻_终极保真版.docx", "磨煤机终极保真版")
